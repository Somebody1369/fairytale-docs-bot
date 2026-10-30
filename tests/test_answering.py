import io
from datetime import timedelta
from zoneinfo import ZoneInfo

import docx
import pytest

from bot import texts
from bot.answering import DOC_CHARS, PROMPT_DOC_CHARS, Answerer
from bot.extract import extract_text
from bot.places import CONIFER, Scope

from .conftest import NOW, FakeLLM, add_animal, add_archive

TZ = ZoneInfo("Europe/Kyiv")


@pytest.fixture
def forest_docs(store, forest):
    add_archive(store, forest, "Членські внески рукавички Під дубом",
                "Внесок 4 жолуді з дорослого звіра на місяць. Здавати до 10 числа.")
    add_archive(store, forest, "Членські внески рукавички У соснах", "Внесок 3 шишки на місяць.")
    add_archive(store, forest, "Як утеплити рукавичку в хвойному лісі",
                "Порада для тих, чия рукавичка стоїть у хвойному лісі. Утеплюйте рукавичку в жовтні.")
    add_archive(store, forest, "Календар сезонів",
                "Цей календар для всіх мешканців лісу. Рукавички утеплюють у листопаді.")
    add_archive(store, forest, "Що робити при пожежі",
                "Ця пам'ятка для всіх мешканців лісу. Довгий безперервний крик означає пожежу. Біжіть до річки.")
    return store


async def test_place_dependent_question_without_place_asks(forest_docs, forest):
    llm = FakeLLM({"needs_place": True, "found": False})
    reply = await Answerer(forest_docs, forest, llm, TZ).ask("скільки внесок?", NOW)
    assert reply.kind == "need_place"
    system, prompt = llm.prompts[0]
    assert "Місце в питанні не вказано" in system
    # The model sees the dues of different mittens side by side and can tell they differ.
    assert "Під дубом" in prompt and "У соснах" in prompt


async def test_model_may_answer_without_place_when_it_says_so(forest_docs, forest):
    llm = FakeLLM({"needs_place": False, "found": True, "answer": "4 жолуді.", "sources": [1]})
    reply = await Answerer(forest_docs, forest, llm, TZ).ask("скільки внесок?", NOW)
    assert reply.kind == "answer"


async def test_whole_forest_question_is_answered_without_place(forest_docs, forest):
    llm = FakeLLM({"found": True, "answer": "Біжіть до річки.", "sources": [1]})
    reply = await Answerer(forest_docs, forest, llm, TZ).ask("що робити при пожежі?", NOW)
    assert reply.kind == "answer"
    assert 'Джерело: "Що робити при пожежі" (Лісовий архів, місце: увесь ліс)' in reply.text


async def test_only_documents_for_the_named_place_reach_the_model(forest_docs, forest):
    llm = FakeLLM({"found": True, "answer": "Утеплюйте в жовтні.", "sources": [1, 2]})
    reply = await Answerer(forest_docs, forest, llm, TZ).ask("коли утеплювати рукавичку у соснах?", NOW)
    prompt = llm.prompts[0][1]
    assert "хвойному лісі" in prompt and "Календар сезонів" in prompt
    assert "Під дубом" not in prompt
    # The more specific (forest part) document goes first.
    assert prompt.index("Як утеплити рукавичку в хвойному лісі") < prompt.index("Календар сезонів")
    assert "Сьогодні: середа, 28 жовтня 2026." in prompt
    assert "Жовтень 2026: 1 чт, 2 пт, 3 сб" in prompt and "Листопад 2026: 1 нд" in prompt
    assert reply.text.startswith("Утеплюйте в жовтні.")
    assert "Джерела:" in reply.text


async def test_mitten_chosen_by_button_is_used(forest_docs, forest):
    llm = FakeLLM({"found": True, "answer": "4 жолуді на місяць.", "sources": [1]})
    reply = await Answerer(forest_docs, forest, llm, TZ).ask("скільки внесок?", NOW, mitten="Під дубом")
    assert "Членські внески рукавички Під дубом" in llm.prompts[0][1]
    assert "У соснах" not in llm.prompts[0][1].split("Документи:")[1]
    assert 'Джерело: "Членські внески рукавички Під дубом" (Лісовий архів, місце: рукавичка "Під дубом")' in reply.text


async def test_model_without_answer_gives_honest_not_found(forest_docs, forest):
    llm = FakeLLM({"found": False, "answer": "", "sources": []})
    reply = await Answerer(forest_docs, forest, llm, TZ).ask("скільки внесок під дубом за малечу?", NOW)
    assert reply.kind == "not_found" and reply.text == texts.NOT_FOUND


async def test_answer_without_sources_is_not_trusted(forest_docs, forest):
    # Two documents could be the source: the bot asks once more, then gives up rather than guess.
    no_sources = {"found": True, "answer": "Щось вигадане.", "sources": []}
    llm = FakeLLM(no_sources, dict(no_sources))
    reply = await Answerer(forest_docs, forest, llm, TZ).ask("коли утеплювати рукавичку у соснах?", NOW)
    assert reply.kind == "not_found" and len(llm.prompts) == 2


async def test_forgotten_source_is_asked_again(forest_docs, forest):
    llm = FakeLLM({"found": True, "answer": "У жовтні.", "sources": []},
                  {"found": True, "answer": "У жовтні.", "sources": ["1"]})
    reply = await Answerer(forest_docs, forest, llm, TZ).ask("коли утеплювати рукавичку у соснах?", NOW)
    assert reply.kind == "answer" and 'Джерело: "Як утеплити рукавичку в хвойному лісі"' in reply.text


async def test_animal_source_shows_author_and_time(store, forest):
    add_animal(store, "Гриби під ялинами", "Під ялинами ростуть рижики після дощу.", Scope("part", CONIFER),
               name="Білочка", when=NOW - timedelta(hours=3))
    llm = FakeLLM({"found": True, "answer": "Рижики після дощу.", "sources": [1]})
    reply = await Answerer(store, forest, llm, TZ).ask("які гриби ростуть під ялиною?", NOW)
    assert ('Джерело: "Гриби під ялинами", надіслано 28.10.2026 о 11:00, автор: Білочка, місце: хвойний ліс'
            in reply.text)
    assert "надіслав звір Білочка" in llm.prompts[0][1]


async def test_no_matching_document_asks_model_to_pick_by_title(forest_docs, forest):
    llm = FakeLLM(picks=[{"ids": []}])
    reply = await Answerer(forest_docs, forest, llm, TZ).ask("як навчитися літати під дубом?", NOW)
    assert reply.kind == "not_found"
    assert len(llm.pick_prompts) == 1 and llm.prompts == []  # title picking only, no answer call


async def test_general_advice_is_not_given_when_places_differ(forest_docs, forest):
    # The forest calendar says November, the conifer forest says October: ask which mitten.
    llm = FakeLLM({"needs_place": True})
    reply = await Answerer(forest_docs, forest, llm, TZ).ask("коли утеплювати рукавичку?", NOW)
    assert reply.kind == "need_place"
    prompt = llm.prompts[0][1]
    assert "Календар сезонів" in prompt and "Як утеплити рукавичку в хвойному лісі" in prompt


async def test_whole_forest_question_does_not_mention_place_rule(forest_docs, forest):
    llm = FakeLLM({"found": True, "answer": "Біжіть до річки.", "sources": [1]})
    await Answerer(forest_docs, forest, llm, TZ).ask("що робити при пожежі?", NOW)
    assert "Місце в питанні не вказано" not in llm.prompts[0][0]


async def test_named_resident_gives_the_place(store, forest):
    add_archive(store, forest, "Статут ОСМД У ліщині", "Мешканці. У рукавичці живуть Ведмідь, Сорока і Вивірка.")
    add_archive(store, forest, "Правила співжиття в рукавичці У ліщині", "З листопада Ведмідь лягає в сплячку.")
    add_archive(store, forest, "Календар сезонів", "Цей календар для всіх мешканців лісу. Хто спить узимку, лягає в грудні.")
    llm = FakeLLM({"found": True, "answer": "З листопада.", "sources": [1]})
    reply = await Answerer(store, forest, llm, TZ).ask("коли Ведмідь лягає спати?", NOW)
    assert reply.kind == "answer"
    # The place comes from the statute before the search, so the model never has to ask for it.
    system, prompt = llm.prompts[0]
    assert 'рукавичка "У ліщині"' in prompt and "за звіром Ведмідь" in prompt
    assert "Місце в питанні не вказано" not in system
    assert "Правила співжиття в рукавичці У ліщині" in prompt


async def test_resident_question_finds_the_mittens_own_documents(store, forest):
    add_archive(store, forest, "Статут ОСМД Під дубом", "У рукавичці живуть Кабан, Сова і Їжак.")
    add_archive(store, forest, "Чергування в рукавичці Під дубом", "Вночі чергує Сова: вона стежить за входом.")
    add_archive(store, forest, "Чергування в рукавичці У соснах", "Вночі вартує Пугач, а вдень Дятел.")
    llm = FakeLLM({"found": True, "answer": "Сова чергує вночі.", "sources": [1]})
    await Answerer(store, forest, llm, TZ).ask("Що робить Сова вночі?", NOW)
    prompt = llm.prompts[0][1].split("Документи:")[1]
    assert "Чергування в рукавичці Під дубом" in prompt and "У соснах" not in prompt


async def test_mitten_advice_is_not_pushed_out_by_stronger_general_matches(store, forest):
    # "Коли збирати жолуді під дубом?": forest and part documents match the words better,
    # yet the mitten's own documents still reach the model.
    add_archive(store, forest, "Коли падають жолуді й горіхи",
                "Порада для листяного лісу. Жолуді збирати у вересні, жолуді сушити, жолуді зберігати.")
    add_archive(store, forest, "Що їсти в листяному лісі", "Порада для листяного лісу. Жолуді, горіхи, жолуді.")
    add_archive(store, forest, "Календар сезонів", "Для всіх мешканців лісу. Жолуді збирати восени, коли жолуді падають.")
    add_archive(store, forest, "Що робити при пожежі", "Для всіх мешканців лісу. Не збирати нічого, біжіть до річки.")
    add_archive(store, forest, "Небезпеки листяного лісу", "Порада для листяного лісу. Жолуді на стежці слизькі.")
    add_archive(store, forest, "Членські внески рукавички Під дубом", "Внесок 4 жолуді на місяць.")
    add_archive(store, forest, "Чергування в рукавичці Під дубом", "Хто чергує, той і збирати жолуді до комори йде.")
    add_archive(store, forest, "Чим поживитись поруч із рукавичкою Під дубом",
                "Жолуді під нашим дубом достигають у другій половині вересня.")
    llm = FakeLLM({"found": True, "answer": "У другій половині вересня.", "sources": [1]})
    await Answerer(store, forest, llm, TZ).ask("Коли збирати жолуді під дубом?", NOW)
    prompt = llm.prompts[0][1]
    assert "Чим поживитись поруч із рукавичкою Під дубом" in prompt
    # The mitten's documents go first: they are the most specific.
    assert prompt.index("Чим поживитись") < prompt.index("Коли падають жолуді")


async def test_title_pick_finds_advice_with_other_words(store, forest):
    # "Що їсти" shares no word with "Чим поживитись": the model picks it by the title.
    add_archive(store, forest, "Чим поживитись поруч із рукавичкою Під вербою", "Верболіз, кора, ягоди калини.")
    add_archive(store, forest, "Що їсти біля річки", "Ця пам'ятка для мешканців біля річки. Що їсти: рогіз, очерет.")
    llm = FakeLLM({"found": True, "answer": "Кора й ягоди калини.", "sources": [1]}, picks=[{"ids": [1, 2]}])
    await Answerer(store, forest, llm, TZ).ask("Що їсти під вербою?", NOW)
    assert "1. Чим поживитись поруч із рукавичкою Під вербою (рукавичка \"Під вербою\")" in llm.pick_prompts[0]
    assert "Чим поживитись поруч із рукавичкою Під вербою" in llm.prompts[0][1]


async def test_two_animal_documents_for_the_same_place_both_reach_the_model(store, forest):
    add_archive(store, forest, "Що робити при пожежі",
                "Для всіх мешканців лісу. При пожежі збираємося біля бобрової греблі. Пожежа: кричати, пожежа.")
    add_animal(store, "Пожежа", "При пожежі збираємося під дубом, пожежа.", Scope("mitten", "Під дубом"),
               when=NOW - timedelta(days=5))
    add_animal(store, "Новий збір", "Тепер при пожежі збираємося на галявині.", Scope("mitten", "Під дубом"),
               author_id=2, name="Сова", when=NOW - timedelta(days=1))
    llm = FakeLLM({"found": True, "answer": "На галявині.", "sources": [1]})
    await Answerer(store, forest, llm, TZ).ask("Де збиратися при пожежі під дубом?", NOW)
    prompt = llm.prompts[0][1]
    assert "Новий збір" in prompt and '"Пожежа"' in prompt
    assert prompt.index("Новий збір") < prompt.index('"Пожежа"')  # newer first


async def test_only_date_questions_get_longer_thinking(forest_docs, forest):
    llm = FakeLLM({"found": True, "answer": "4 жолуді.", "sources": [1]},
                  {"found": True, "answer": "До 10 числа.", "sources": [1]})
    answerer = Answerer(forest_docs, forest, llm, TZ)
    await answerer.ask("скільки внесок під дубом?", NOW)
    await answerer.ask("до якого числа здавати внесок під дубом?", NOW)
    assert llm.reasoning == ["low", "medium"]


def test_newer_archive_document_wins_among_equals():
    from bot.answering import ANSWER_SYSTEM
    assert "Серед документів архіву, як і серед документів від звірів, важливіший новіший" in ANSWER_SYSTEM


async def test_animal_documents_of_one_day_carry_the_time(store, forest):
    add_animal(store, "Збір", "При пожежі збираємося під дубом.", Scope("mitten", "Під дубом"),
               when=NOW - timedelta(hours=3))
    add_animal(store, "Новий збір", "При пожежі тепер збираємося на галявині.", Scope("mitten", "Під дубом"),
               author_id=2, name="Сова", when=NOW - timedelta(hours=1))
    llm = FakeLLM({"found": True, "answer": "На галявині.", "sources": [1]})
    await Answerer(store, forest, llm, TZ).ask("Де збиратися при пожежі під дубом?", NOW)
    prompt = llm.prompts[0][1]
    assert "надіслав звір Сова, 28 жовтня 2026 о 13:00" in prompt and "надіслав звір Зайчик, 28 жовтня 2026 о 11:00" in prompt


async def test_guide_section_weighs_as_the_whole_guide(store, forest):
    # Mishka: a section about "Під дубом" in a whole-forest guide is still whole-forest advice,
    # so the deciduous forest's own advice comes first. A model that puts the guide first is asked again.
    add_archive(store, forest, "Жолуді листяного лісу", "Порада для листяного лісу. Жолуді збирайте в жовтні.")
    add_animal(store, "Путівник лісом", "Путівник для всіх звірів.\n\nПід дубом. Жолуді збирайте у вересні.",
               Scope("forest"), name="Сорока")
    llm = FakeLLM({"found": True, "answer": "У вересні.", "sources": [2, 1]},
                  {"found": True, "answer": "У жовтні.", "sources": [1, 2]})
    reply = await Answerer(store, forest, llm, TZ).ask("коли збирати жолуді під дубом?", NOW)
    first, second = llm.prompts[0][1], llm.prompts[1][1]
    assert "Рівень: частина лісу" in first and "Рівень: увесь ліс" in first
    assert 'документ [1] (рівень "частина лісу") конкретніший за документ [2] (рівень "увесь ліс")' in second
    assert reply.text.startswith("У жовтні.") and reply.text.index("Жолуді листяного") < reply.text.index("Путівник")


def word_guide(sections: list[tuple[str, str]], intro: str = "") -> str:
    """A guide as Word gives it: paragraphs joined by single line breaks, no blank lines."""
    document = docx.Document()
    if intro:
        document.add_paragraph(intro)
    for heading, body in sections:
        document.add_paragraph(heading)
        for _ in range(3):
            document.add_paragraph(body)
    buffer = io.BytesIO()
    document.save(buffer)
    return extract_text(buffer.getvalue(), "Путівник.docx")


FILLER = "Восени дні коротшають, тож запаси робіть заздалегідь і перевіряйте їх щотижня після дощу. "


async def test_long_word_guide_reaches_the_model(store, forest):
    # Word separates paragraphs with a single line break: the whole guide used to be one chunk
    # over the size limit, and the model got the guide's title with an empty text.
    text = word_guide([("Запаси на зиму", FILLER * 4), ("Гриби", FILLER * 4),
                       ("Жолуді", "Під дубом запасайте п'ять кошиків жолудів. " + FILLER)])
    assert len(text) > DOC_CHARS and "\n\n" not in text
    add_animal(store, "Зимовий путівник", text, Scope("mitten", "Під дубом"), name="Сова")
    llm = FakeLLM({"found": True, "answer": "П'ять кошиків.", "sources": [1]})
    await Answerer(store, forest, llm, TZ).ask("скільки кошиків жолудів запасати під дубом?", NOW)
    excerpt = llm.prompts[0][1].split("Текст:\n", 1)[1]
    assert "п'ять кошиків жолудів" in excerpt and len(excerpt) <= DOC_CHARS + 50


async def test_long_guide_gives_the_section_of_the_asked_place(store, forest):
    # "біля сосон" shares no word with the heading "У соснах"; the forest map still finds the section.
    sections = [(f"{place}", f"{fact} {FILLER * 2}") for place, fact in [
        ("Під ялиною", "Під ялиною ростуть рижики."), ("Під дубом", "Біля дуба шукайте білі гриби."),
        ("У ліщині", "У ліщині ростуть лисички."), ("У соснах", "Маслюки ростуть на галявинах."),
        ("Біля бобрової греблі", "Біля греблі ростуть опеньки."), ("На узліссі", "На узліссі ростуть печериці.")]]
    add_animal(store, "Грибний путівник", word_guide(sections, "Путівник для всіх мешканців лісу."), Scope("forest"))
    llm = FakeLLM({"found": True, "answer": "Маслюки.", "sources": [1]})
    await Answerer(store, forest, llm, TZ).ask("які гриби біля сосон?", NOW)
    excerpt = llm.prompts[0][1].split("Текст:\n", 1)[1]
    assert "У соснах\nМаслюки ростуть на галявинах." in excerpt  # the heading stays with its section


def document_texts(prompt: str) -> list[str]:
    return [block.split("Текст:\n", 1)[1] for block in prompt.split("\n\n[")[1:]]


@pytest.mark.parametrize("lengths", [[40 * len(FILLER)] * 6, [2400] * 7, [2400] * 5 + [6000] * 2])
async def test_documents_together_fit_one_prompt(store, forest, lengths):
    # Groq's free tier rejects a prompt over about 8000 tokens, so long and middling documents share a cut.
    for i, length in enumerate(lengths):
        add_animal(store, f"Поради про жолуді {i}", ("Під дубом жолуді. " + FILLER * 70)[:length],
                   Scope("mitten", "Під дубом"), author_id=i, when=NOW - timedelta(hours=i))
    picks = list(range(len(lengths) - 3, len(lengths) + 1))  # the title pick and keyword search take different ones
    llm = FakeLLM({"found": True, "answer": "x", "sources": [1]}, picks=[{"ids": picks}])
    await Answerer(store, forest, llm, TZ).ask("Коли запасати жолуді під дубом?", NOW)
    texts = document_texts(llm.prompts[0][1])
    assert len(texts) >= 5 and all(texts) and sum(map(len, texts)) <= PROMPT_DOC_CHARS


LATE = " Це важливо, бо зима довга, а весна приходить пізно."


def word_text(paragraphs: list[str]) -> str:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return extract_text(buffer.getvalue(), "Путівник.docx")


WINTER_GUIDE = [
    "Зимовий путівник для всіх мешканців лісу",
    "Тут зібрано поради, як перезимувати в кожній рукавичці. Знайдіть розділ своєї рукавички й прочитайте його уважно.",
    "Під ялиною",
    "Під ялиною тепло навіть у мороз, бо гілки низько. Вхід утеплюйте мохом, а не листям: листя мокне." + LATE,
    "Запаси тримайте в сухій хвої біля коріння, шишки перевіряйте щотижня, бо волога їх псує, а миші люблять сухі." + LATE,
    "Вечорами не розпалюйте вогню: хвоя легко займається. Світіть світлячками, їх приносить Сорока." + LATE,
    "У соснах",
    "У соснах вітряно, тож рукавичку прив'язуйте до коріння лозою. Перевіряйте вузли після кожної бурі." + LATE,
    "Смолу не чіпайте: вона потрібна для латання рукавички, і латає її лише Дятел." + LATE,
    "Маслюки сушіть на гілках до першого снігу, потім знімайте, бо замерзнуть." + LATE,
    "Біля мурашника",
    "Біля мурашника не копайте ям для запасів: мурахи сердяться й кусають навіть узимку, коли прокидаються у відлигу." + LATE,
    "Запаси тримайте на високому пеньку, накривайте корою від снігу." + LATE,
    "Стежкою повз мурашник ходіть лише вдень, щоб не наступити на мурашині ходи." + LATE,
    "Під дубом",
    "Під дубом взимку найбільше гостей, бо тут найбільша комора лісу. Гостей пускайте лише вдень." + LATE,
    "Жолуді зберігайте в коморі під корінням, а не в дуплі, бо там їх знаходять сойки й розносять по лісу." + LATE,
    "Кожен мешканець приносить до комори п'ять кошиків жолудів до першого снігу, а Кабан приносить сім. "
    "Хто не приніс, той чергує біля комори тиждень." + LATE,
    "У ліщині",
    "У ліщині горіхи сушіть три дні на сонці, тоді вони не гниють у коморі й не гірчать." + LATE,
    "Шкаралупу не викидайте біля входу, бо її бачать лисиці й приходять до рукавички вночі." + LATE,
    "Ведмідь лягає спати в листопаді, тож його частку горіхів відкладайте заздалегідь." + LATE,
    "Під вербою",
    "Під вербою вода восени піднімається, тож запаси тримайте на гілках, а не на землі." + LATE,
    "Утеплюйте рукавичку очеретом, його носять Жабка й Качка з річки щосуботи." + LATE,
    "Кору верби не обдирайте взимку: дерево замерзне й навесні не дасть листя." + LATE,
    "Біля бобрової греблі",
    "Біля греблі не ходіть по тонкому льоду, особливо в грудні, коли лід ще не встиг стати." + LATE,
    "Бобри відкривають греблю в березні, до того всі запаси мають бути на високому березі." + LATE,
    "Рибу морозьте в ополонці, яку рубає Бобер щоранку." + LATE,
    "На узліссі",
    "На узліссі вітер холодний, тож заносьте запаси глибше до лісу, до старого пня." + LATE,
    "Мисливці взимку ходять узліссям, тож вхід ховайте за кущем шипшини." + LATE,
    "Шипшину сушіть і зберігайте в сухій норі, вона знадобиться від застуди." + LATE,
]


async def guide_excerpt(store, forest, paragraphs: list[str], question: str) -> str:
    add_animal(store, "Путівник", word_text(paragraphs), Scope("forest"), name="Сова")
    llm = FakeLLM({"found": True, "answer": "x", "sources": [1]})
    await Answerer(store, forest, llm, TZ).ask(question, NOW)
    return document_texts(llm.prompts[0][1])[0]


async def test_end_of_the_asked_section_is_kept(store, forest):
    # A chunk may hold the end of the asked section and then the next heading: that end still belongs to it.
    excerpt = await guide_excerpt(store, forest, WINTER_GUIDE, "Що треба робити взимку під дубом?")
    assert "п'ять кошиків жолудів" in excerpt


async def test_section_of_the_asked_forest_part_is_kept(store, forest):
    more = LATE * 2
    paragraphs = ["Як утеплити рукавичку", "Поради для всіх мешканців лісу на зиму." + more,
                  "Під ялинами", "Утеплюйте рукавичку хвоєю й мохом, шар у три лапи." + more,
                  "Вхід закривайте ялиновою гілкою." + more, "Хвою міняйте раз на місяць." + more,
                  "Серед дубів і ліщини", "Утеплюйте рукавичку сухим дубовим листям, шар у дві лапи." + more,
                  "Листя збирайте лише сухе, після трьох сонячних днів." + more, "Мокре листя гниє й холодить." + more,
                  "Біля річки", "Утеплюйте рукавичку очеретом, а не мохом, і піднімайте її на палі." + more,
                  "Очерет ріжте в жовтні, поки він сухий." + more, "Палі ставте до першого льоду." + more,
                  "На узліссі", "Утеплюйте рукавичку сіном і ставте заслін від вітру." + more,
                  "Сіно беріть у полі після покосу." + more, "Заслін плетіть із лози." + more]
    excerpt = await guide_excerpt(store, forest, paragraphs, "Як утеплити рукавичку біля річки?")
    assert "Палі ставте до першого льоду" in excerpt


async def test_section_after_the_asked_one_does_not_take_its_place(store, forest):
    # The hazel section after a short oak section is not oak advice; the deciduous forest's advice is.
    more = LATE + " І запаси мають дожити до квітня."
    hazel = ["Шкаралупу не викидайте біля входу, бо її бачать лисиці й приходять уночі.",
             "Ведмідь лягає спати в листопаді, тож його частку горіхів відкладайте раніше.",
             "Кошики для горіхів плетіть із лози заздалегідь, великі кошики не беріть.",
             "Вхід закривайте гілками щовечора, бо вітер холодний.",
             "Гостей із сусідніх рукавичок пускайте лише вдень і записуйте на корі.",
             "Старі кошики латайте лозою, нові плетіть лише в морозні дні.",
             "Мішечки підвішуйте до стелі, щоб миші не дісталися.", "Сушені ягоди тримайте в глечиках із кришкою.",
             "Лозу для кошиків ріжте восени біля струмка.", "Горіхи для Білочки складайте окремо, у малий кошик.",
             "Глечики перевіряйте раз на тиждень, чи не відсиріли.", "Підлогу застеляйте сухим мохом."]
    deciduous = ["Узимку в листяному лісі стежки розчищає Кабан.", "Сніг біля входу відкидайте вранці, поки він м'який.",
                 "Листя під деревами не згрібайте: воно гріє коріння.",
                 "Комору в листяному лісі відкривають лише по суботах, ключ від комори тримає Сова.",
                 "Хворим звірам видають без черги.", "Сухі гілки збирайте лише з землі.",
                 "Воду беріть лише з джерела, не з калюж.", "Дітей самих далеко не відпускайте."]
    paragraphs = (["Зимовий путівник", "Поради на зиму для мешканців рукавичок." + more,
                   "Під дубом", "Під дубом запасайте п'ять кошиків жолудів на кожного мешканця до першого снігу." + more,
                   "У ліщині", "У ліщині горіхи сушіть три дні на сонці, тоді вони не гниють у коморі." + more]
                  + [line + more for line in hazel] + ["Листяний ліс загалом"] + [line + more for line in deciduous])
    excerpt = await guide_excerpt(store, forest, paragraphs, "Коли відкривають комору під дубом?")
    assert "п'ять кошиків жолудів" in excerpt and "по суботах" in excerpt


def excerpt_of(store, forest, lines: list[str], question: str, mittens: list[str], parts: list[str]) -> str:
    add_animal(store, "Путівник", "\n".join(lines), Scope("forest"))
    doc = store.active_documents()[0]
    assert len(doc.text) > DOC_CHARS
    return Answerer(store, forest, FakeLLM(), TZ)._excerpt(doc, question, mittens, parts, DOC_CHARS)


NOISY = "Гриби ростуть після дощу, гриби ростуть у затінку, гриби ростуть рядами. "


def test_asked_place_beats_word_overlap(store, forest):
    sections = [("У листяному лісі", NOISY * 3)] * 3 + [("Під ялиною", NOISY * 3)] * 2 + \
               [("Під дубом", "Білі шукайте біля коріння старого дерева. " + FILLER * 2)]
    text = excerpt_of(store, forest, [line for heading, body in sections for line in (heading, body, body, body)],
                      "які гриби ростуть під дубом?", ["Під дубом"], [])
    assert "Білі шукайте біля коріння" in text


def test_section_running_over_chunks_keeps_its_place(store, forest):
    lines = []
    for heading in ["Під ялиною", "Під дубом", "У ліщині"]:
        lines += [heading] + [NOISY * 3] * 2
    lines += ["У соснах", FILLER * 7, FILLER * 7, "Маслюки збирайте на галявинах."]
    for heading in ["На узліссі", "Під вербою", "Біля бобрової греблі"]:
        lines += [heading] + [NOISY * 3] * 2
    text = excerpt_of(store, forest, lines, "які гриби ростуть біля сосон?", ["У соснах"], [])
    assert "Маслюки збирайте" in text


def test_part_question_takes_the_sections_of_its_mittens(store, forest):
    lines = []
    for mitten in ["Під дубом", "У ліщині", "На узліссі", "Біля бобрової греблі", "Під вербою", "Під ялиною", "У соснах"]:
        lines += [mitten] + [f"Тут живуть звірі рукавички {mitten}. {FILLER * 3}"] * 3
    text = excerpt_of(store, forest, lines, "що робити у хвойному лісі?", [], ["conifer"])
    assert "рукавички Під ялиною" in text and "рукавички У соснах" in text


def test_excerpt_finds_synonyms(store, forest):
    lines = []
    for topic in ["Запаси", "Прибирання", "Сусіди", "Свята", "Нічна тиша", "Чим поживитись"]:
        lines += [topic] + [FILLER * 4] * 3 + (["Кора і бруньки дуже поживні."] if topic == "Чим поживитись" else [])
    assert "Кора і бруньки" in excerpt_of(store, forest, lines, "що їсти?", [], [])


async def test_part_question_about_mitten_advice_asks_which_mitten(store, forest):
    # Dues are written per mitten. "біля річки" names a forest part, so the river's mittens are offered,
    # instead of "no such advice".
    add_archive(store, forest, "Членські внески рукавички Під вербою", "Внесок 3 горіхи на місяць.")
    add_archive(store, forest, "Членські внески рукавички Біля бобрової греблі", "Внесок 2 рибини на місяць.")
    add_archive(store, forest, "Членські внески рукавички Під дубом", "Внесок 4 жолуді на місяць.")
    llm = FakeLLM({"needs_place": True})
    reply = await Answerer(store, forest, llm, TZ).ask("Скільки внесок біля річки?", NOW)
    assert reply.kind == "need_place" and reply.mittens == ["Під вербою", "Біля бобрової греблі"]
    assert reply.text == texts.ASK_MITTEN_IN_PART.format(part="біля річки")
    system, prompt = llm.prompts[0]
    assert "У питанні названо частину лісу (біля річки)" in system
    assert "Під вербою" in prompt and "Біля бобрової греблі" in prompt and "4 жолуді" not in prompt


async def test_part_question_answers_from_the_parts_own_advice_first(store, forest):
    add_archive(store, forest, "Що їсти біля річки", "Ця пам'ятка для мешканців біля річки. Їжте рогіз і очерет.")
    add_archive(store, forest, "Чим поживитись поруч із рукавичкою Під вербою", "Їжте кору верби.")
    llm = FakeLLM({"found": True, "answer": "Рогіз і очерет; під вербою ще кору.", "sources": [1, 2]})
    reply = await Answerer(store, forest, llm, TZ).ask("Що їсти біля річки?", NOW)
    prompt = llm.prompts[0][1]
    assert prompt.index("Що їсти біля річки") < prompt.index("Чим поживитись")
    # The mitten's advice is an exception to the part's: the model is not asked to put it first.
    assert len(llm.prompts) == 1 and reply.text.index("Що їсти біля річки") < reply.text.index("Чим поживитись")


async def test_terminal_ask_works_with_empty_settings(monkeypatch, capsys):
    # docker-compose passes unset settings as empty strings: TZ_NAME="" is Kyiv, not a crash.
    from bot import ask
    from bot.answering import Reply

    used = {}

    class Answering:
        def __init__(self, store, forest, llm, tz):
            used["tz"] = tz

        async def ask(self, question, now, mitten=None):
            return Reply("answer", "4 жолуді.")

    async def no_sync(*args):
        used["ocr"] = args[2].ocr_langs

    for name in ("TZ_NAME", "OCR_LANGS"):
        monkeypatch.setenv(name, "")
    monkeypatch.setenv("DB_PATH", ":memory:")
    monkeypatch.setenv("GROQ_API_KEY", "x")
    monkeypatch.setattr(ask, "sync_archive", no_sync)
    monkeypatch.setattr(ask, "Answerer", Answering)
    await ask.main("скільки внесок під дубом?")
    assert str(used["tz"]) == "Europe/Kyiv" and used["ocr"] == "ukr"
    assert capsys.readouterr().out.strip() == "4 жолуді."


def test_text_without_blank_lines_is_cut_into_chunks():
    from bot.search import CHUNK_CHARS, split_chunks
    one_line = FILLER * 30  # a TXT guide written as one long line
    lines = "\n".join(["Під дубом"] + [FILLER * 2] * 20)  # Word: one line per paragraph
    ending = lines + "\nЖолуді\nЗбирайте до 10 жовтня"  # last lines without a full stop are kept too
    for text in (one_line, lines, ending):
        chunks = split_chunks(text)
        assert len(chunks) > 1 and all(len(chunk) <= CHUNK_CHARS + 81 for chunk in chunks)
        assert "".join(chunks).replace("\n", "").replace(" ", "") == text.replace("\n", "").replace(" ", "")
    assert split_chunks("Календар сезонів\nУтеплюйте в листопаді.") == ["Календар сезонів\nУтеплюйте в листопаді."]


async def test_part_question_answered_with_a_mitten_first_is_asked_again(store, forest):
    # Rule 3 says a mitten's advice comes first; with a forest part asked it is only an exception.
    add_archive(store, forest, "Що їсти біля річки", "Ця пам'ятка для мешканців біля річки. Їжте рогіз і очерет.")
    add_archive(store, forest, "Чим поживитись поруч із рукавичкою Під вербою", "Їжте кору верби.")
    llm = FakeLLM({"found": True, "answer": "Під вербою кору.", "sources": [2, 1]},
                  {"found": True, "answer": "Рогіз і очерет; під вербою ще кору.", "sources": [1, 2]})
    reply = await Answerer(store, forest, llm, TZ).ask("Що їсти біля річки?", NOW)
    assert "на відміну від правила 3" in llm.prompts[0][0]
    assert "у питанні названо частину лісу, а документ [2] лише для однієї рукавички" in llm.prompts[1][1]
    assert reply.text.startswith("Рогіз і очерет")
    # Only the mitten's advice answers: nothing to put before it.
    llm = FakeLLM({"found": True, "answer": "Під вербою кору.", "sources": [2]})
    reply = await Answerer(store, forest, llm, TZ).ask("Що їсти біля річки?", NOW)
    assert len(llm.prompts) == 1 and reply.kind == "answer"


async def test_part_question_keeps_the_whole_forest_advice(store, forest):
    # Many mittens' documents must not push the part's and the whole forest's own advice out of the prompt.
    add_archive(store, forest, "Що їсти біля річки", "Ця пам'ятка для мешканців біля річки. Що їсти: рогіз.")
    add_archive(store, forest, "Що їсти взимку", "Для всіх мешканців лісу. Що їсти взимку: жолуді й горіхи.")
    for mitten in ("Під вербою", "Біля бобрової греблі"):
        for i in range(3):
            add_animal(store, f"Що їсти {mitten} {i}", f"Що їсти: кору, ягоди, їжу {i}.", Scope("mitten", mitten),
                       author_id=i + 1)
    llm = FakeLLM({"found": True, "answer": "Рогіз.", "sources": [1]})
    await Answerer(store, forest, llm, TZ).ask("Що їсти біля річки?", NOW)
    prompt = llm.prompts[0][1]
    assert "Що їсти біля річки" in prompt and "Що їсти взимку" in prompt


OAK = WINTER_GUIDE.index("Під дубом")


@pytest.mark.parametrize("start", ["До ліщини по горіхи ходіть лише вдень, а жолуді",
                                   "Ліщиновими гілками закривайте вхід щовечора, а жолуді"])
async def test_paragraph_naming_another_place_does_not_end_the_section(store, forest, start):
    guide = list(WINTER_GUIDE)
    guide[OAK + 2] = start + " зберігайте в коморі під корінням, а не в дуплі." + LATE
    excerpt = await guide_excerpt(store, forest, guide, "Що треба робити взимку під дубом?")
    assert "п'ять кошиків жолудів" in excerpt


async def test_wrapped_pdf_line_naming_another_place_does_not_end_the_section(store, forest):
    import textwrap
    guide = list(WINTER_GUIDE)
    guide[OAK + 1] = ("Під дубом взимку найбільше гостей, бо тут найбільша комора лісу. Своїх гостей пускайте лише "
                      "вдень, а сусідів із ліщини та з-під верби можна пускати й увечері." + LATE)
    add_animal(store, "Путівник", "\n".join(line for p in guide for line in textwrap.wrap(p, 76)), Scope("forest"))
    llm = FakeLLM({"found": True, "answer": "x", "sources": [1]})
    await Answerer(store, forest, llm, TZ).ask("Що треба робити взимку під дубом?", NOW)
    assert "Хто не приніс, той чергує біля комори тиждень" in " ".join(document_texts(llm.prompts[0][1])[0].split())


async def test_place_named_in_passing_does_not_outrank_the_asked_section(store, forest):
    guide = list(WINTER_GUIDE)
    for heading, line in [("Під ялиною", "Шишки міняйте на жолуді в коморі під дубом, але не частіше разу на місяць."),
                          ("Під вербою", "Жолуді для малечі просіть у сусідів під дубом заздалегідь."),
                          ("На узліссі", "Шипшину носіть до комори під дубом, там її сушать для всіх.")]:
        guide[guide.index(heading) + 2] = line + LATE
    excerpt = await guide_excerpt(store, forest, guide, "Що треба робити взимку під дубом?")
    assert "п'ять кошиків жолудів" in excerpt and "Жолуді зберігайте в коморі" in excerpt


async def test_no_place_asked_is_not_turned_into_one_mittens_answer(store, forest):
    # The model answers a whole-forest question from the whole-forest memo [2] and mentions the mitten's
    # exception [1] after it: with no place asked, nothing makes it put the mitten first.
    add_archive(store, forest, "Що робити при пожежі", "Ця пам'ятка для всіх мешканців лісу. При пожежі біжіть до річки.")
    add_archive(store, forest, "Небезпеки поруч із рукавичкою Під дубом", "При пожежі біжіть до дуба.")
    llm = FakeLLM({"found": True, "answer": "Біжіть до річки; під дубом до дуба.", "sources": [2, 1]})
    await Answerer(store, forest, llm, TZ).ask("що робити при пожежі?", NOW)
    assert "Небезпеки поруч" in llm.prompts[0][1].split("[1]")[1].split("[2]")[0] and len(llm.prompts) == 1
