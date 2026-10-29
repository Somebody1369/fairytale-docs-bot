"""Text extraction from PDF, Word (.docx and the old .doc) and plain text files.

Some archive PDFs are scans on a birch-bark background with no text layer,
so pages without text are rendered and passed through Tesseract OCR.
"""

import io
import logging
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

log = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt"}
MIN_PAGE_TEXT = 20
MAX_OCR_PAGES = 30  # a scanned page takes a few seconds to read; a longer scan is refused rather than cut


class TooManyScannedPages(ValueError):
    def __init__(self, pages: int):
        super().__init__(f"{pages} scanned pages, at most {MAX_OCR_PAGES} are read")
        self.pages = pages


class UnsupportedFormat(Exception):
    pass


def extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def title_from_filename(filename: str) -> str:
    stem = Path(filename).stem
    stem = re.sub(r"\s*\(\d+\)$", "", stem)  # "Назва (1).pdf" is the same document
    return re.sub(r"[_\s]+", " ", stem).strip()


def old_word_supported() -> bool:
    """The old Word format is read by antiword (installed in the Docker image)."""
    return shutil.which("antiword") is not None


def extract_text(data: bytes, filename: str, ocr_langs: str = "ukr", ext: str | None = None) -> str:
    ext = ext or extension(filename)
    if ext == ".txt":
        return _clean(_decode_text(data))
    if ext == ".docx":
        return _clean(_docx_text(data))
    if ext == ".doc":
        return _clean(_doc_text(data))
    if ext == ".pdf":
        return _clean(_pdf_text(data, ocr_langs))
    raise UnsupportedFormat(ext)


def _decode_text(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _docx_text(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _doc_text(data: bytes) -> str:
    if not old_word_supported():
        raise UnsupportedFormat(".doc")
    with tempfile.NamedTemporaryFile(suffix=".doc") as file:
        file.write(data)
        file.flush()
        result = subprocess.run(["antiword", "-m", "UTF-8.txt", "-w", "0", file.name],
                                capture_output=True, timeout=60, check=True)
    return result.stdout.decode("utf-8", errors="replace")


def _pdf_text(data: bytes, ocr_langs: str) -> str:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    pages = [pdf[index].get_textpage().get_text_range() for index in range(len(pdf))]
    scanned = [index for index, text in enumerate(pages) if len(text.strip()) < MIN_PAGE_TEXT]
    if len(scanned) > MAX_OCR_PAGES:
        raise TooManyScannedPages(len(scanned))
    for index in scanned:
        pages[index] = _ocr_page(pdf[index], ocr_langs) or pages[index]
    return "\n\n".join(pages)


def _ocr_page(page, ocr_langs: str) -> str:
    try:
        import pytesseract
        from PIL import ImageOps
    except ImportError:
        log.warning("OCR is not available, scanned page skipped")
        return ""

    image = page.render(scale=300 / 72).to_pil().convert("L")
    image = ImageOps.autocontrast(image)
    # Text is near-black; the bark background and its stains are lighter.
    image = image.point(lambda value: 0 if value < 110 else 255)
    try:
        raw = pytesseract.image_to_string(image, lang=ocr_langs, config="--psm 4")
    except pytesseract.TesseractError as error:
        log.warning("OCR failed: %s", error)
        return ""
    return _clean_ocr(raw)


def _clean_ocr(raw: str) -> str:
    """Drop specks of the bark background that OCR reads as stray symbols."""
    lines = []
    for line in raw.splitlines():
        line = re.sub(r"^[^\w«\"(]+", "", line).strip()
        # OCR reads the digit 3 as the Cyrillic letter З: "становить З шишки", "1 рибина за З гілки".
        line = re.sub(r"(?<=[а-яіїєґ0-9,:] )З(?= )", "3", line)
        if len(re.findall(r"[А-Яа-яІіЇїЄєҐґA-Za-z]", line)) >= 3:
            lines.append(line)
        elif not line and lines and lines[-1]:
            lines.append("")
    return "\n".join(lines)


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("­", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
